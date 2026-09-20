"""Upload flow tests: text extraction, dedicated resume/JD priority, cache reset.

No live LLM/API calls. Everything runs against a tmp_path knowledge dir, so the
real data/ folder is never touched.
"""
import os
import sys

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

from dotenv import load_dotenv

load_dotenv(os.path.join(BASE_DIR, ".env"))

import agent  # noqa: E402
import app  # noqa: E402
import pytest  # noqa: E402


def test_extract_text_md_and_txt(tmp_path):
    path = tmp_path / "简历.md"
    path.write_text("我的简历内容", encoding="utf-8")

    assert app.extract_text(str(path), "简历.md") == "我的简历内容"
    assert app.extract_text("直接传字节".encode(), "简历.txt") == "直接传字节"


def test_extract_docx_includes_tables(tmp_path):
    """Resumes are usually laid out as tables; paragraphs alone would drop them."""
    import docx

    document = docx.Document()
    document.add_paragraph("个人简介段落")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "技能"
    table.rows[0].cells[1].text = "Python / LangChain"
    path = tmp_path / "resume.docx"
    document.save(str(path))

    text = app.extract_text(str(path), "resume.docx")
    assert "个人简介段落" in text
    assert "Python / LangChain" in text


def test_extract_pdf(tmp_path):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.cell(text="Resume PDF with magic number 24680.")
    path = tmp_path / "resume.pdf"
    pdf.output(str(path))

    assert "24680" in app.extract_text(str(path), "resume.pdf")


def test_extract_empty_pdf_raises(tmp_path):
    """An image-only (scanned) PDF yields no text — the UI must be able to warn."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    path = tmp_path / "scanned.pdf"
    pdf.output(str(path))

    with pytest.raises(ValueError, match="没有提取到文字"):
        app.extract_text(str(path), "scanned.pdf")


def test_extract_unsupported_ext():
    with pytest.raises(ValueError, match="不支持的文件格式"):
        app.extract_text(b"x", "resume.pages")


def test_uploaded_resume_takes_priority(tmp_path):
    (tmp_path / "论文全文.txt").write_text("论文正文", encoding="utf-8")
    agent.save_uploaded_doc("上传的简历", "resume", data_dir=str(tmp_path))

    assert agent.load_resume(str(tmp_path)) == "上传的简历"


def test_resume_falls_back_to_data_dir_when_nothing_uploaded(tmp_path):
    (tmp_path / "简历项目经历.md").write_text("旧简历", encoding="utf-8")

    assert agent.load_resume(str(tmp_path)) == "旧简历"


def test_uploaded_jd_takes_priority(tmp_path):
    (tmp_path / "目标JD-旧.md").write_text("旧JD", encoding="utf-8")
    agent.save_uploaded_doc("新 JD 文本", "jd", data_dir=str(tmp_path))

    assert agent.load_jd(str(tmp_path)) == "新 JD 文本"


def test_uploaded_resume_does_not_leak_into_jd_slot(tmp_path):
    agent.save_uploaded_doc("简历内容", "resume", data_dir=str(tmp_path))

    assert agent.load_jd(str(tmp_path)) is None


def test_save_uploaded_doc_rejects_unknown_kind():
    with pytest.raises(ValueError, match="unknown kind"):
        agent.save_uploaded_doc("x", "cover_letter")


def test_reset_checker_drops_cached_pipeline():
    """Stale Chroma handles would query the collection a rebuild just deleted."""
    agent._checker = ("fake-store",)
    agent.reset_checker()

    assert agent._checker is None
